from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.resume.engine import ResumeIntelligenceEngine


def make_pdf(lines: list[str], columns: bool = False) -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=letter)
    if columns:
        for index, text in enumerate(lines):
            column = index % 2
            row = index // 2
            pdf.drawString(40 + column * 270, 740 - row * 24, text)
    else:
        for index, text in enumerate(lines):
            pdf.drawString(40, 740 - index * 24, text)
    pdf.save()
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Taylor Morgan")
    document.add_paragraph("taylor@example.com | +1 555 123 4567 | Austin, TX")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Backend engineer building reliable services.")
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Platform Engineer | Acme | Jan 2022 - Present")
    document.add_paragraph("- Built FastAPI services with Python and PostgreSQL for 10,000 users.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.S. Computer Science, State University, 2021")
    document.add_paragraph("SKILLS")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ReactJS"
    table.rows[0].cells[1].text = "Node.js"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_extracts_sections_entities_skills_and_evidence() -> None:
    content = make_pdf([
        "Taylor Morgan",
        "taylor@example.com | linkedin.com/in/taylor | Austin, TX",
        "SUMMARY",
        "Backend engineer building reliable services.",
        "EXPERIENCE",
        "Platform Engineer | Acme | Jan 2022 - Present",
        "- Built FastAPI services with Python and PostgreSQL for 10,000 users.",
        "EDUCATION",
        "B.S. Computer Science, State University, 2021",
        "SKILLS",
        "Python, FastAPI, PostgreSQL, ReactJS, Node.js",
        "PROJECTS",
        "Resume Parser",
        "- Built a parser using Python and FastAPI.",
    ])

    extraction, text = ResumeIntelligenceEngine().parse_bytes(content, "pdf")

    assert "Taylor Morgan" in text
    assert extraction.candidate.name is not None
    assert extraction.contact.email is not None
    assert extraction.contact.linkedin is not None
    assert extraction.sections[0].name == "summary"
    assert extraction.experience
    assert extraction.experience[0].technologies == ["Python", "FastAPI", "PostgreSQL"]
    assert {skill.name for skill in extraction.skills["Frameworks"]} >= {"FastAPI", "React", "Node.js"}
    assert extraction.projects[0].measurable_result_present is False
    assert extraction.projects[0].uncertain_fields == ["impact"]
    assert any(item.page_number == 1 for item in extraction.evidence)


def test_docx_table_and_aliases_are_supported() -> None:
    extraction, _ = ResumeIntelligenceEngine().parse_bytes(make_docx(), "docx")

    assert extraction.document_signals is not None
    assert extraction.document_signals.table_count == 1
    framework_names = {skill.name for skill in extraction.skills["Frameworks"]}
    assert {"React", "Node.js"}.issubset(framework_names)
    assert extraction.contact.phone is not None
    assert extraction.experience[0].duration == "ongoing (exact duration uncertain)"


def test_column_like_pdf_and_missing_sections_are_handled_conservatively() -> None:
    extraction, _ = ResumeIntelligenceEngine().parse_bytes(make_pdf(["Alex Example", "alex@example.com", "EXPERIENCE", "Volunteer", "- Helped customers"], columns=True), "pdf")

    assert extraction.contact.email is not None
    assert extraction.experience
    missing_codes = {signal.code for signal in extraction.quality_signals}
    assert "missing_education" in missing_codes
    assert "missing_skills" in missing_codes
    assert any(item.weak for item in extraction.bullet_analysis)


def test_nontechnical_resume_does_not_receive_invented_skills() -> None:
    extraction, _ = ResumeIntelligenceEngine().parse_bytes(make_pdf(["Morgan Lee", "morgan@example.com", "EXPERIENCE", "Office Coordinator", "- Coordinated schedules and welcomed visitors.", "EDUCATION", "High School Diploma, 2020"]), "pdf")

    assert extraction.skills == {}
    assert extraction.experience[0].responsibilities == ["Coordinated schedules and welcomed visitors."]
