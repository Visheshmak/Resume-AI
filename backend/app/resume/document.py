from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class DocumentLine:
    text: str
    page_number: int
    line_number: int
    bold: bool = False
    style_name: str | None = None
    is_bullet: bool = False


@dataclass
class ExtractedDocument:
    lines: list[DocumentLine] = field(default_factory=list)
    page_count: int = 0
    table_count: int = 0
    source_format: str = ""
    notes: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)


def extract_document(content: bytes, file_type: str) -> ExtractedDocument:
    if file_type.lower() == "pdf":
        return _extract_pdf(content)
    if file_type.lower() == "docx":
        return _extract_docx(content)
    raise ValueError(f"Unsupported document type: {file_type}")


def _extract_pdf(content: bytes) -> ExtractedDocument:
    reader = PdfReader(BytesIO(content))
    result = ExtractedDocument(page_count=len(reader.pages), source_format="pdf")
    line_number = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            page_text = page.extract_text() or ""
        page_lines = page_text.splitlines()
        for raw_line in page_lines:
            text = " ".join(raw_line.split())
            if not text:
                continue
            line_number += 1
            result.lines.append(DocumentLine(text, page_number, line_number, is_bullet=_is_bullet(text)))
    if not result.lines:
        result.notes.append("No extractable text was found; the document may be image-based and require OCR.")
    return result


def _extract_docx(content: bytes) -> ExtractedDocument:
    try:
        document = Document(BytesIO(content))
    except (AttributeError, KeyError, OSError, ValueError) as exc:
        raise ValueError("Invalid DOCX document structure") from exc
    result = ExtractedDocument(page_count=1, table_count=len(document.tables), source_format="docx")
    line_number = 0

    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        line_number += 1
        bold = any(run.bold for run in paragraph.runs if run.text.strip())
        result.lines.append(
            DocumentLine(
                text,
                1,
                line_number,
                bold=bool(bold),
                style_name=paragraph.style.name if paragraph.style else None,
                is_bullet=_is_bullet(text) or (paragraph.style.name or "").lower().startswith("list bullet"),
            )
        )

    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            for text in filter(None, cells):
                line_number += 1
                result.lines.append(DocumentLine(text, 1, line_number, is_bullet=_is_bullet(text)))
        result.notes.append("Table content was flattened into reading order; original cell coordinates are not retained.")
    return result


def _is_bullet(text: str) -> bool:
    return text[:1] in {"-", "*", "•", "‣", "▪", "◦"} or text[:2] in {"- ", "* "}
